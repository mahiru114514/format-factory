import importlib.util
import io
import json
import time
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from PIL import Image


def test_service_is_available():
    assert importlib.util.find_spec('mobile_server.app') is not None


@pytest.fixture
def service(tmp_path):
    from mobile_server.app import create_app
    from mobile_server.config import Settings
    app = create_app(Settings(storage_dir=tmp_path / 'jobs', cleanup_interval_seconds=.05))
    with TestClient(app) as client:
        yield client, app.state.jobs, tmp_path / 'jobs'


def png():
    output = io.BytesIO()
    Image.new('RGB', (8, 8), 'red').save(output, 'PNG')
    return output.getvalue()


def submit(client, name='test.png', content=None, category='images', target='jpg', **extra):
    data = {'category': category, 'target': target, 'consent': 'true', 'options': '{}'}
    data.update(extra)
    return client.post('/v1/jobs', files={'file': (name, png() if content is None else content)}, data=data)


def auth(job):
    return {'Authorization': 'Bearer ' + job['token']}


def finished(client, job, timeout=30):
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        response = client.get('/v1/jobs/' + job['id'], headers=auth(job))
        assert response.status_code == 200, response.text
        if response.json()['state'] in ('succeeded', 'failed'):
            return response.json()
        time.sleep(.05)
    pytest.fail('Job did not finish in time')


def result_files(client, job):
    status = finished(client, job)
    assert status['state'] == 'succeeded', status
    response = client.get('/v1/jobs/' + job['id'] + '/result', headers=auth(job))
    assert response.status_code == 200
    assert response.headers['content-type'] == 'application/zip'
    assert status['result_name'].endswith('.zip')
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def test_capabilities_are_truthful(service):
    client, _, _ = service
    payload = client.get('/v1/capabilities').json()
    assert payload['max_upload_bytes'] == 50 * 1024 * 1024
    assert payload['worker_timeout_seconds'] == 120
    assert payload['retention_seconds'] == 3600
    assert payload['result_format'] == 'zip'
    assert set(payload['categories']) == {'images', 'tables', 'documents', 'text'}
    assert payload['engines']['libreoffice']['ready'] == bool(payload['engines']['libreoffice']['path'])
    assert all('ready' in pair for pair in payload['categories']['documents']['pairs'])


@pytest.mark.parametrize('consent', ['', 'false', 'yes', '1'])
def test_no_consent_never_stores_input(service, consent):
    client, _, storage = service
    response = submit(client, consent=consent)
    assert response.status_code == 403
    assert list(storage.iterdir()) == []


def test_missing_consent_is_forbidden(service):
    client, _, storage = service
    response = client.post('/v1/jobs', files={'file': ('test.png', png())},
                           data={'category': 'images', 'target': 'jpg'})
    assert response.status_code == 403
    assert list(storage.iterdir()) == []


@pytest.mark.parametrize('name,content,category,target', [
    ('a.png', b'not a PNG', 'images', 'jpg'),
    ('a.docx', png(), 'documents', 'txt'),
    ('a.csv', png(), 'tables', 'xlsx'),
    ('a.exe', b'text', 'documents', 'pdf'),
    ('a.png', png(), 'images', 'exe'),
    ('a.png', png(), 'documents', 'txt'),
    ('../../escaped.png', png(), 'images', 'jpg'),
    ('C:\\private\\escaped.png', png(), 'images', 'jpg'),
])
def test_invalid_type_format_and_names_rejected_before_storage(service, name, content, category, target):
    client, _, storage = service
    response = submit(client, name, content, category, target)
    assert response.status_code in (400, 415, 422), response.text
    assert not list(storage.iterdir())


def test_bounded_upload(tmp_path):
    from mobile_server.app import create_app
    from mobile_server.config import Settings
    with TestClient(create_app(Settings(storage_dir=tmp_path / 'jobs', max_upload_bytes=32))) as client:
        assert submit(client).status_code == 413
    assert not list((tmp_path / 'jobs').iterdir())


def test_status_download_delete_and_bearer_isolation(service):
    client, _, storage = service
    one = submit(client).json()
    two = submit(client).json()
    for suffix, method in [('', 'get'), ('/result', 'get'), ('', 'delete')]:
        url = '/v1/jobs/' + one['id'] + suffix
        assert getattr(client, method)(url).status_code == 401
        assert getattr(client, method)(url, headers=auth(two)).status_code == 404
    files = result_files(client, one)
    assert list(files) == ['test.jpg']
    with Image.open(io.BytesIO(files['test.jpg'])) as image:
        assert image.format == 'JPEG'
    assert client.delete('/v1/jobs/' + one['id'], headers=auth(one)).status_code == 204
    assert client.get('/v1/jobs/' + one['id'], headers=auth(one)).status_code == 404
    assert not (storage / one['id']).exists()


def test_expiry_removes_input_output_and_token(service):
    client, manager, storage = service
    job = submit(client).json()
    finished(client, job)
    manager.cleanup(now=time.time() + 3601)
    assert not (storage / job['id']).exists()
    assert client.get('/v1/jobs/' + job['id'], headers=auth(job)).status_code == 404


def test_table_real_round_trip_preserves_literal_formula(service):
    client, _, _ = service
    job = submit(client, 'data.csv', '名称,编号\n北京,001\n公式,=1+2\n'.encode(), 'tables', 'xlsx').json()
    data = result_files(client, job)['data.xlsx']
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(data))
    assert book.active['B2'].value == '001'
    assert book.active['B3'].value == '=1+2'
    assert book.active['B3'].data_type == 's'
    book.close()
    reverse = submit(client, 'data.xlsx', data, 'tables', 'csv').json()
    assert '北京,001' in result_files(client, reverse)['data.csv'].decode('utf-8-sig')


def docx_bytes(image=False):
    from docx import Document
    document = Document()
    document.add_paragraph('Portable document')
    if image:
        document.add_picture(io.BytesIO(png()))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def rewritten_zip(content, name, data):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(output, 'w') as archive:
        for item in source.infolist():
            if item.filename != name:
                archive.writestr(item.filename, source.read(item))
        archive.writestr(name, data)
    return output.getvalue()


@pytest.mark.parametrize('name,data', [
    ('../../escape', b'danger'),
    ('word/_rels/document.xml.rels', b'''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId9" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="http://127.0.0.1/private" TargetMode="External"/></Relationships>'''),
    ('word/_rels/document.xml.rels', b'''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId9" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="file:///etc/passwd"/></Relationships>'''),
])
def test_docx_rejects_zip_traversal_external_resources(service, name, data):
    client, _, storage = service
    response = submit(client, 'a.docx', rewritten_zip(docx_bytes(), name, data), 'documents', 'md')
    assert response.status_code == 400, response.text
    assert not list(storage.iterdir())


def test_docx_rejects_decompressed_bomb(service):
    client, _, _ = service
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', '<Types/>')
        archive.writestr('word/document.xml', b'x' * (16 * 1024 * 1024))
    response = submit(client, 'bomb.docx', output.getvalue(), 'documents', 'txt')
    assert response.status_code == 400, response.text


@pytest.mark.parametrize('text', [
    '![secret](file:///etc/passwd)', '![remote](http://127.0.0.1/private)',
    '<img src="http://127.0.0.1/private">', '\\input{/etc/passwd}',
    '![x][ref]\n\n[ref]: ../../secret.png',
])
def test_markdown_resource_inclusions_rejected(service, text):
    client, _, _ = service
    response = submit(client, 'a.md', text.encode(), 'documents', 'docx')
    assert response.status_code == 400, response.text


def test_docx_markdown_keeps_packaged_media(service):
    client, _, _ = service
    job = submit(client, 'a.docx', docx_bytes(image=True), 'documents', 'md').json()
    files = result_files(client, job)
    assert 'Portable document' in files['a.md'].decode()
    media = [name for name in files if name.startswith('media/')]
    assert media
    assert media[0] in files['a.md'].decode()


def test_markdown_docx_and_txt_real_conversion(service):
    client, _, _ = service
    job = submit(client, 'a.md', b'# Hello\n\nWorld', 'documents', 'docx').json()
    docx = result_files(client, job)['a.docx']
    job = submit(client, 'a.docx', docx, 'documents', 'txt').json()
    assert 'Hello' in result_files(client, job)['a.txt'].decode('utf-8')


def test_pdf_text_page_zero_is_used(service):
    import pymupdf
    document = pymupdf.open()
    document.new_page().insert_text((40, 40), 'First page')
    document.new_page().insert_text((40, 40), 'Second page')
    data = document.tobytes()
    document.close()
    source = service[2].parent / 'two.pdf'
    source.write_bytes(data)
    # PDF document extraction can avoid OCR; text category is OCR/formula only.
    job = submit(service[0], 'two.pdf', data, 'documents', 'txt', options='{"page":0}').json()
    text = result_files(service[0], job)['two.txt'].decode()
    assert 'First page' in text and 'Second page' not in text


def test_text_options_reject_unknown_method_negative_page_and_non_txt(service):
    client, _, _ = service
    for options, target in [(' {"method":"bogus"}', 'txt'), ('{"page":-1}', 'txt'), ('{}', 'pdf')]:
        response = submit(client, category='text', target=target, options=options)
        assert response.status_code == 400, response.text


def test_missing_libreoffice_is_actionable(service):
    client, _, _ = service
    capabilities = client.get('/v1/capabilities').json()
    if capabilities['engines']['libreoffice']['ready']:
        pytest.skip('LibreOffice is installed; covered by optional PDF integration')
    response = submit(client, 'a.docx', docx_bytes(), 'documents', 'pdf')
    assert response.status_code == 503
    assert 'LibreOffice' in response.json()['detail']


def test_real_worker_timeout_kills_nested_child(tmp_path):
    import os
    import sys
    import psutil
    from mobile_server.manager import run_isolated
    pid_file = tmp_path / 'child.pid'
    code = "import subprocess,sys,time,pathlib; p=subprocess.Popen([sys.executable,'-c','import time;time.sleep(60)']); pathlib.Path(sys.argv[1]).write_text(str(p.pid));time.sleep(60)"
    started = time.monotonic()
    with pytest.raises(TimeoutError):
        run_isolated([sys.executable, '-c', code, str(pid_file)], timeout=.6, cwd=tmp_path)
    assert time.monotonic() - started < 6
    pid = int(pid_file.read_text())
    for _ in range(30):
        if not psutil.pid_exists(pid):
            break
        time.sleep(.05)
    assert not psutil.pid_exists(pid) or psutil.Process(pid).status() == psutil.STATUS_ZOMBIE


def test_worker_queue_is_bounded_and_cleanup_is_automatic(tmp_path):
    from mobile_server.app import create_app
    from mobile_server.config import Settings
    settings = Settings(storage_dir=tmp_path / 'jobs', max_jobs=1, retention_seconds=60,
                        cleanup_interval_seconds=.05)
    application = create_app(settings)
    with TestClient(application) as client:
        job = submit(client).json()
        assert submit(client).status_code == 429
        # Isolated validation can outlast a tiny real-time retention window.
        # Check capacity first, then age the job and observe the janitor itself.
        with application.state.jobs.lock:
            application.state.jobs.jobs[job['id']].created_at = time.time() - 61
        end = time.monotonic() + 5
        while (settings.storage_dir / job['id']).exists() and time.monotonic() < end:
            time.sleep(.05)
        assert not (settings.storage_dir / job['id']).exists()
        assert client.get('/v1/jobs/' + job['id'], headers=auth(job)).status_code == 404


def test_docx_dynamic_fields_cannot_fetch_local_files(service):
    client, _, storage = service
    content = docx_bytes()
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        xml = archive.read('word/document.xml')
    xml = xml.replace(b'</w:body>', b'<w:p><w:fldSimple w:instr="INCLUDETEXT &quot;file:///etc/passwd&quot;"/></w:p></w:body>')
    response = submit(client, 'a.docx', rewritten_zip(content, 'word/document.xml', xml), 'documents', 'pdf')
    assert response.status_code == 400
    assert not list(storage.iterdir())


def complex_word_field(*instructions):
    runs = b''.join(b'<w:r><w:instrText>' + text + b'</w:instrText></w:r>' for text in instructions)
    return (b'<w:r><w:fldChar w:fldCharType="begin"/></w:r>' + runs +
            b'<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>value</w:t></w:r>'
            b'<w:r><w:fldChar w:fldCharType="end"/></w:r>')


@pytest.mark.parametrize('fields', [
    b'<w:fldSimple w:instr="DATE"/><w:fldSimple w:instr="INCLUDETEXT &quot;file:///etc/passwd&quot;"/>',
    complex_word_field(b'DATE') + complex_word_field(b'INCLUDETEXT &quot;file:///etc/passwd&quot;'),
    complex_word_field(b'DATE') + complex_word_field(b'INCLU', b'DETEXT &quot;file:///etc/passwd&quot;'),
    complex_word_field(b'INCLU', b'DETEXT &quot;file:///etc/passwd&quot;'),
], ids=['simple-after-date', 'complex-after-date', 'split-after-date', 'split-alone'])
def test_docx_checks_each_field_and_reconstructs_split_instruction_runs(service, fields):
    content = docx_bytes()
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        xml = archive.read('word/document.xml')
    xml = xml.replace(b'</w:body>', b'<w:p>' + fields + b'</w:p></w:body>')
    response = submit(service[0], 'fields.docx', rewritten_zip(content, 'word/document.xml', xml), 'documents', 'md')
    assert response.status_code == 400, response.text
    assert not list(service[2].iterdir())


def test_multi_sheet_zip_uses_original_name_for_all_outputs(service):
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.title = 'First'
    workbook.active.append(['one'])
    workbook.create_sheet('Second').append(['two'])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    job = submit(service[0], 'book.xlsx', output.getvalue(), 'tables', 'csv').json()
    files = result_files(service[0], job)
    assert set(files) == {'book_First.csv', 'book_Second.csv'}


@pytest.mark.parametrize('method,example,expected', [
    ('ocr', '文字识别示例.png', '123'),
    ('formula', '公式识别示例.png', '='),
])
def test_real_recognition_runs_in_isolated_service_worker(service, method, example, expected):
    sample = Path(__file__).resolve().parents[2] / '示例文件' / example
    job = submit(service[0], example, sample.read_bytes(), 'text', 'txt',
                 options=json.dumps({'method': method, 'page': 0})).json()
    files = result_files(service[0], job)
    assert expected in next(iter(files.values())).decode('utf-8')


def test_real_pdf_ocr_uses_zero_based_selected_page(service):
    import pymupdf
    sample = Path(__file__).resolve().parents[2] / '示例文件' / '文字识别示例.png'
    document = pymupdf.open()
    first = document.new_page(width=750, height=150)
    first.insert_image(first.rect, stream=sample.read_bytes())
    document.new_page(width=750, height=150)
    data = document.tobytes()
    document.close()
    job = submit(service[0], 'scan.pdf', data, 'text', 'txt', options='{"method":"ocr","page":0}').json()
    assert '123' in result_files(service[0], job)['scan.txt'].decode('utf-8')


def test_failed_conversion_is_reported_without_success_or_result(service):
    job = submit(service[0], category='text', target='txt').json()  # blank red image has no text
    assert finished(service[0], job)['state'] == 'failed'
    assert service[0].get('/v1/jobs/' + job['id'] + '/result', headers=auth(job)).status_code == 409


def test_real_pdf_output_when_libreoffice_installed(service):
    client, _, _ = service
    if not client.get('/v1/capabilities').json()['engines']['libreoffice']['ready']:
        pytest.skip('LibreOffice not installed on this host')
    job = submit(client, 'a.docx', docx_bytes(image=True), 'documents', 'pdf').json()
    assert result_files(client, job)['a.pdf'].startswith(b'%PDF-')


def test_timeout_reports_failed_job_and_keeps_result_unavailable(tmp_path):
    from mobile_server.app import create_app
    from mobile_server.config import Settings
    with TestClient(create_app(Settings(storage_dir=tmp_path / 'jobs', worker_timeout_seconds=.001))) as client:
        job = submit(client).json()
        status = finished(client, job)
        assert status['state'] == 'failed'
        assert '超时' in status['error']
        assert client.get('/v1/jobs/' + job['id'] + '/result', headers=auth(job)).status_code == 409


def test_hyperlinks_are_preserved_without_being_fetched(service):
    client, _, _ = service
    content = docx_bytes()
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        xml = archive.read('word/_rels/document.xml.rels')
    xml = xml.replace(b'</Relationships>', b'<Relationship Id="remote-link" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.invalid" TargetMode="External"/></Relationships>')
    response = submit(client, 'linked.docx', rewritten_zip(content, 'word/_rels/document.xml.rels', xml), 'documents', 'txt')
    assert response.status_code == 202
    assert 'Portable document' in result_files(client, response.json())['linked.txt'].decode()


def test_restart_discards_only_owned_orphan_jobs(tmp_path):
    from mobile_server.app import create_app
    from mobile_server.config import Settings
    root = tmp_path / 'jobs'
    owned = root / ('a' * 32)
    unrelated = root / ('b' * 32)
    owned.mkdir(parents=True)
    unrelated.mkdir()
    (owned / 'request.json').write_text('{"owner":"mobile_server_v1"}')
    (unrelated / 'keep.txt').write_text('keep')
    with TestClient(create_app(Settings(storage_dir=root))):
        assert not owned.exists()
        assert (unrelated / 'keep.txt').read_text() == 'keep'


def test_chunked_oversized_body_stops_without_storage(service):
    client, manager, storage = service
    maximum = manager.settings.max_upload_bytes + 65536
    def body():
        for _ in range(maximum // 65536 + 2):
            yield b'x' * 65536
    response = client.post('/v1/jobs', content=body(), headers={'content-type': 'multipart/form-data; boundary=x'})
    assert response.status_code == 413
    assert not list(storage.iterdir())


def test_configuration_rejects_unbounded_worker_values(tmp_path):
    from mobile_server.config import Settings
    with pytest.raises(ValueError):
        Settings(storage_dir=tmp_path, max_workers=0)


@pytest.mark.skipif(__import__('os').name != 'nt', reason='Windows console launcher')
def test_windows_libreoffice_uses_console_launcher(tmp_path, monkeypatch):
    from mobile_server.engines import libreoffice_path
    executable = tmp_path / 'soffice.exe'
    console = tmp_path / 'soffice.com'
    executable.touch()
    console.touch()
    monkeypatch.setenv('FACTORY_SOFFICE', str(executable))
    assert libreoffice_path() == str(console)


def test_ocr_readiness_requires_packaged_model_files(tmp_path, monkeypatch):
    from mobile_server import engines
    from types import SimpleNamespace
    original = engines.importlib.util.find_spec
    def lookup(name, *args, **kwargs):
        if name == 'rapidocr_onnxruntime':
            return SimpleNamespace(origin=str(tmp_path / '__init__.py'))
        return original(name, *args, **kwargs)
    monkeypatch.setattr(engines.importlib.util, 'find_spec', lookup)
    assert not engines.engine_status()['ocr']['ready']


@pytest.mark.parametrize('target', ['md', 'docx'])
def test_document_respects_explicit_encoding_for_utf16le_without_bom(service, target):
    text = '你好世界，保留指定编码。'
    job = submit(service[0], 'encoded.txt', text.encode('utf-16-le'), 'documents', target,
                 options='{"encoding":"utf-16"}').json()
    files = result_files(service[0], job)
    if target == 'md':
        assert text in files['encoded.md'].decode('utf-8')
    else:
        from docx import Document
        document = Document(io.BytesIO(files['encoded.docx']))
        assert text in '\n'.join(p.text for p in document.paragraphs)


@pytest.mark.parametrize('encoding,codec', [('gb18030', 'gb18030'), ('utf-16', 'utf-16-le')])
def test_markdown_is_normalized_from_explicit_encoding_before_pandoc(service, encoding, codec):
    from docx import Document
    text = '中文标题与指定编码'
    job = submit(service[0], 'encoded.md', ('# ' + text).encode(codec), 'documents', 'docx',
                 options=json.dumps({'encoding': encoding})).json()
    document = Document(io.BytesIO(result_files(service[0], job)['encoded.docx']))
    assert text in '\n'.join(p.text for p in document.paragraphs)


def test_deep_content_validation_does_not_run_in_api_process(service, monkeypatch):
    from mobile_server import validation
    api = __import__('mobile_server.app', fromlist=['create_app'])
    def forbidden(*args, **kwargs):
        raise AssertionError('Untrusted parser ran inside API process')
    monkeypatch.setattr(validation, 'validate_content', forbidden)
    monkeypatch.setattr(api, 'validate_content', forbidden, raising=False)
    # Patching the API interpreter cannot affect the separately launched worker.
    response = submit(service[0])
    assert response.status_code == 202, response.text
    assert result_files(service[0], response.json())['test.jpg'].startswith(b'\xff\xd8')


def test_deep_validation_timeout_is_bounded_and_cleans_staging(tmp_path):
    from mobile_server.app import create_app
    from mobile_server.config import Settings
    storage = tmp_path / 'jobs'
    app = create_app(Settings(storage_dir=storage, validation_timeout_seconds=.001))
    started = time.monotonic()
    with TestClient(app) as client:
        response = submit(client)
        assert response.status_code == 408, response.text
        assert '校验超时' in response.json()['detail']
        assert not list(storage.iterdir())
    assert time.monotonic() - started < 6


def test_deep_validation_worker_has_process_memory_limit(service, monkeypatch):
    from mobile_server import manager
    original = manager.run_isolated
    calls = []
    def isolated(*args, **kwargs):
        calls.append(kwargs.get('memory_limit_bytes'))
        return original(*args, **kwargs)
    monkeypatch.setattr(manager, 'run_isolated', isolated)
    response = submit(service[0], content=b'not an image')
    assert response.status_code == 400
    assert calls == [768 * 1024 * 1024]
    assert not list(service[2].iterdir())
