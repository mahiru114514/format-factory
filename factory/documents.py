from __future__ import annotations

import re
import subprocess
from html import escape
from pathlib import Path
from tempfile import TemporaryDirectory

from .common import ConversionResult, Options, normalized_format, publish, read_text


def pandoc(source: Path, target: str, destination: Path | None = None, input_format: str | None = None) -> str:
    import pypandoc
    source = Path(source).resolve()
    destination = Path(destination).resolve() if destination else None
    command = [pypandoc.get_pandoc_path(), str(source), '--to', target, '--wrap=none', '--resource-path', str(source.parent)]
    if input_format:
        command += ['--from', input_format]
    if destination:
        command += ['--output', str(destination)]
    result = subprocess.run(command, cwd=str(source.parent), capture_output=True, timeout=120,
                            creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
    if result.returncode:
        raise ValueError('文档转换失败：' + result.stderr.decode('utf-8', errors='replace')[:1200])
    return result.stdout.decode('utf-8')


def literal_markdown(text: str) -> str:
    return '\n'.join(re.sub(r'([\\`*_{}\[\]()#+.!|>~$<-])', r'\\\1', line) + '  ' for line in text.splitlines())


def extract_pdf(source: Path, ocr: bool = False, page: int | None = None) -> str:
    import pymupdf
    output = []
    with pymupdf.open(source) as document:
        if document.needs_pass:
            raise ValueError('PDF 有密码保护，请先在阅读器中解密并另存。')
        indices = range(len(document)) if page is None else [page]
        for index in indices:
            if index < 0 or index >= len(document):
                raise ValueError('PDF 页码超出范围。')
            pdf_page = document[index]
            text = pdf_page.get_text(sort=True).strip()
            if ocr:
                from .recognition import recognize_image
                blocks = [(block[1], block[0], block[4].strip()) for block in pdf_page.get_text('blocks')
                          if block[6] == 0 and block[4].strip()]
                regions = [pymupdf.Rect(image['bbox']) & pdf_page.rect for image in pdf_page.get_image_info()]
                if not regions and not text:
                    regions = [pdf_page.rect]
                known_lines = {re.sub(r'\s+', '', line).casefold() for line in text.splitlines()}
                seen_regions = set()
                for region in regions:
                    key = tuple(round(value, 1) for value in region)
                    if region.is_empty or key in seen_regions:
                        continue
                    seen_regions.add(key)
                    pixmap = pdf_page.get_pixmap(matrix=pymupdf.Matrix(2, 2), clip=region, alpha=False)
                    try:
                        recognized = recognize_image(pixmap.tobytes('png'))
                    except ValueError as error:
                        if '未识别到' not in str(error):
                            raise
                        continue
                    lines = []
                    for line in recognized.splitlines():
                        key_line = re.sub(r'\s+', '', line).casefold()
                        if key_line and key_line not in known_lines:
                            lines.append(line)
                            known_lines.add(key_line)
                    if lines:
                        blocks.append((region.y0, region.x0, '\n'.join(lines)))
                text = '\n\n'.join(block[2] for block in sorted(blocks, key=lambda item: (item[0], item[1])))
            elif not text:
                raise ValueError(f'第 {index + 1} 页没有可提取的文字，请开启“识别扫描 PDF”。')
            output.append(text)
    result = '\n\n'.join(output)
    if not result.strip():
        raise ValueError('未识别到文字，请检查 PDF 内容或选择更清晰的扫描件。')
    return result


def docx_plain(source: Path) -> str:
    return pandoc(source, 'plain', input_format='docx')


def style_docx(path: Path) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document(path)
    for name in ('Normal', 'Body Text', 'First Paragraph', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        if name in document.styles:
            style = document.styles[name]
            style.font.name = 'Microsoft YaHei'
            rpr = style.element.get_or_add_rPr()
            fonts = rpr.find(qn('w:rFonts'))
            if fonts is None:
                fonts = OxmlElement('w:rFonts')
                rpr.append(fonts)
            fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    document.save(path)


def word_pdf(source: Path, destination: Path) -> None:
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    application = document = None
    try:
        application = win32com.client.DispatchEx('Word.Application')
        application.Visible = False
        application.DisplayAlerts = 0
        application.AutomationSecurity = 3
        document = application.Documents.Open(str(source.resolve()), ReadOnly=True, AddToRecentFiles=False)
        document.ExportAsFixedFormat(str(destination.resolve()), 17, OpenAfterExport=False)
    finally:
        if document is not None:
            document.Close(False)
        if application is not None:
            application.Quit()
        pythoncom.CoUninitialize()


def basic_pdf(text: str, destination: Path) -> None:
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    font_path = Path('C:/Windows/Fonts/simsun.ttc')
    if font_path.exists():
        font_name = 'FactoryChinese'
        if font_name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(font_name, str(font_path), subfontIndex=0))
    else:
        font_name = 'STSong-Light'
        if font_name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    style = ParagraphStyle('Body', fontName=font_name, fontSize=11, leading=18, wordWrap='CJK')
    story = []
    for line in text.splitlines():
        if line.strip():
            story.append(Paragraph(escape(line).replace('  ', '&nbsp;&nbsp;'), style))
        else:
            story.append(Spacer(1, 9))
    if not story:
        story.append(Spacer(1, 1))
    SimpleDocTemplate(str(destination), pagesize=(595.28, 841.89), rightMargin=48, leftMargin=48,
                      topMargin=48, bottomMargin=48).build(story)


def pdf_word(source: Path, destination: Path, ocr: bool) -> list[str]:
    import pymupdf
    with pymupdf.open(source) as document:
        if document.needs_pass:
            raise ValueError('PDF 有密码保护，请先解密并另存。')
        has_scan = any(not page.get_text().strip() or (ocr and page.get_image_info()) for page in document)
    if has_scan:
        text = extract_pdf(source, ocr=ocr)
        from docx import Document
        document = Document()
        for line in text.splitlines():
            document.add_paragraph(line)
        document.save(destination)
        style_docx(destination)
        return ['扫描或含图片的 PDF 已按可编辑文字导出，原始图片与版式未保留；请核对识别结果。']
    from pdf2docx import Converter
    converter = Converter(str(source))
    try:
        converter.convert(str(destination), multi_processing=False)
    finally:
        converter.close()
    return ['PDF 转 Word 后请核对多栏、表格和公式等复杂内容。']


def word_markdown(source: Path, staged: Path, output: Path) -> Path | None:
    """Extract images to a collision-safe folder and keep all links relative."""
    import pypandoc
    import shutil
    from .common import unique_path
    while True:
        media_target = unique_path(output / f'{source.stem}_media')
        try:
            media_target.mkdir()
            break
        except FileExistsError:
            continue
    media_staged = staged.parent / media_target.name
    try:
        command = [pypandoc.get_pandoc_path(), str(source), '-t', 'markdown', '--wrap=none',
                   '--extract-media', media_target.name, '-o', str(staged)]
        result = subprocess.run(command, cwd=str(staged.parent), capture_output=True, timeout=120,
                                creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        if result.returncode:
            raise ValueError('Word 转 Markdown 失败：' + result.stderr.decode('utf-8', errors='replace')[:800])
        if media_staged.exists():
            shutil.copytree(media_staged, media_target, dirs_exist_ok=True)
            return media_target
        media_target.rmdir()
        return None
    except BaseException:
        # This directory was exclusively reserved above and contains only this task's output.
        shutil.rmtree(media_target)
        raise


def convert_document(source: Path, target: str, output: Path, options: Options) -> ConversionResult:
    source, output = Path(source).resolve(), Path(output).resolve()
    source_format = normalized_format(source.suffix)
    if target not in ('txt', 'md', 'docx', 'pdf'):
        raise ValueError('不支持的文档输出格式。')
    output.mkdir(parents=True, exist_ok=True)
    warnings = []
    media_destination = None
    with TemporaryDirectory(prefix='.factory-', dir=output) as temp:
        temp = Path(temp)
        staged = temp / f'result.{target}'
        if source_format == 'pdf':
            if target == 'docx':
                warnings.extend(pdf_word(source, staged, options.pdf_ocr))
            else:
                text = extract_pdf(source, ocr=options.pdf_ocr)
                staged.write_text(literal_markdown(text) if target == 'md' else text, encoding='utf-8')
                if target == 'md':
                    warnings.append('PDF 转 Markdown 以提取文字为主，不自动恢复标题和表格结构。')
        else:
            if source_format == 'txt':
                text = read_text(source)
                prepared = temp / 'input.md'
                prepared.write_text(literal_markdown(text), encoding='utf-8')
                input_format = 'markdown'
            elif source_format == 'md':
                prepared = source
                input_format = 'markdown'
            else:
                prepared = source
                input_format = 'docx'
            if target == 'txt':
                staged.write_text(pandoc(prepared, 'plain', input_format=input_format), encoding='utf-8')
            elif target == 'md':
                if source_format == 'txt':
                    staged.write_text(literal_markdown(text), encoding='utf-8')
                else:
                    media_destination = word_markdown(source, staged, output)
                    if media_destination:
                        warnings.append(f'文档图片保存在 {media_destination.name}，移动 Markdown 时请一并保留该文件夹。')
            else:
                docx = staged if target == 'docx' else temp / 'intermediate.docx'
                if source_format == 'docx' and target == 'pdf':
                    docx = source
                else:
                    pandoc(prepared, 'docx', docx, input_format=input_format)
                    style_docx(docx)
                if target == 'pdf':
                    if options.pdf_engine != 'basic':
                        try:
                            word_pdf(docx, staged)
                        except Exception as error:
                            if options.pdf_engine == 'word':
                                raise ValueError('Word 导出 PDF 失败，请检查 Word 是否可正常启动。') from error
                            basic_pdf(docx_plain(docx), staged)
                            warnings.append('Word 引擎不可用，已使用基础排版导出 PDF；请检查图片、表格和公式。')
                    else:
                        basic_pdf(docx_plain(docx), staged)
                        warnings.append('使用基础文字排版，未保留图片及复杂版式。')
        try:
            destination = publish(staged, output / f'{source.stem}.{target}')
        except BaseException:
            if media_destination:
                import shutil
                shutil.rmtree(media_destination)
            raise
    return ConversionResult([destination], warnings)
