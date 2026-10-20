import csv
import io
from datetime import date, datetime, time
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

from .common import ConversionResult, Options, publish, read_text, safe_name


def convert_table(source: Path, target: str, output: Path, options: Options) -> ConversionResult:
    source, output = Path(source), Path(output)
    if len(options.delimiter) != 1:
        raise ValueError('CSV 分隔符必须是单个字符。')
    output.mkdir(parents=True, exist_ok=True)
    warnings, staged_outputs, published = [], [], []
    with TemporaryDirectory(prefix='.factory-', dir=output) as temp:
        if source.suffix.lower() == '.xlsx' and target == 'csv':
            values = load_workbook(source, read_only=True, data_only=True)
            formulas = load_workbook(source, read_only=True, data_only=False)
            try:
                for index, (sheet, formula_sheet) in enumerate(zip(values.worksheets, formulas.worksheets)):
                    name = source.stem if len(values.worksheets) == 1 else f'{source.stem}_{safe_name(sheet.title)}'
                    staged = Path(temp) / f'{index}.csv'
                    missing = 0
                    with staged.open('w', encoding=options.encoding, newline='') as stream:
                        writer = csv.writer(stream, delimiter=options.delimiter)
                        for row, formula_row in zip(sheet.iter_rows(), formula_sheet.iter_rows()):
                            data = []
                            for cell, original in zip(row, formula_row):
                                if original.data_type == 'f' and cell.value is None:
                                    missing += 1
                                value = cell.value
                                if isinstance(value, datetime):
                                    value = value.isoformat(sep=' ')
                                elif isinstance(value, (date, time)):
                                    value = value.isoformat()
                                data.append(value)
                            writer.writerow(data)
                    if missing:
                        warnings.append(f'工作表“{sheet.title}”有 {missing} 个公式没有缓存值，已输出为空；请用 Excel 重新计算并保存后重试。')
                    staged_outputs.append((staged, output / f'{name}.csv'))
            finally:
                values.close()
                formulas.close()
            warnings.append('CSV 仅保存单张表的数据值，不保留样式、合并单元格和公式本身。')
        elif source.suffix.lower() == '.csv' and target == 'xlsx':
            text = read_text(source, options.encoding)
            book = Workbook(write_only=True)
            sheet = book.create_sheet('数据')
            try:
                for row_index, row in enumerate(csv.reader(io.StringIO(text, newline=''), delimiter=options.delimiter), 1):
                    if row_index > 1_048_576 or len(row) > 16_384:
                        raise ValueError('CSV 超出 Excel 的行数或列数限制。')
                    cells = []
                    for value in row:
                        if len(value) > 32767:
                            raise ValueError(f'第 {row_index} 行存在超过 Excel 单元格长度限制的内容。')
                        cell = WriteOnlyCell(sheet, value=value)
                        cell.data_type = 's'
                        cells.append(cell)
                    sheet.append(cells)
                staged = Path(temp) / 'data.xlsx'
                book.save(staged)
            finally:
                book.close()
            staged_outputs.append((staged, output / f'{source.stem}.xlsx'))
            warnings.append('CSV 内容按文本导入，保留前导零、长编号及以等号开头的原文。')
        else:
            raise ValueError('表格支持 XLSX → CSV 或 CSV → XLSX。')
        try:
            for staged, requested in staged_outputs:
                published.append(publish(staged, requested))
        except BaseException:
            for path in published:
                path.unlink(missing_ok=True)
            raise
    return ConversionResult(published, warnings)
