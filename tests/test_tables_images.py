import csv

import pytest
from PIL import Image
from openpyxl import Workbook, load_workbook


def test_multiple_worksheets_to_csv(tmp_path):
    from factory.tables import convert_table
    from factory.common import Options
    book = Workbook()
    book.active.title = '招生'
    book.active.append(['姓名', '编号', '备注'])
    book.active.append(['张三', '00123', '包含,逗号\n和换行'])
    book.create_sheet('第二张').append(['你好', 123])
    source = tmp_path / '中文.xlsx'
    book.save(source)
    result = convert_table(source, 'csv', tmp_path, Options())
    assert len(result.paths) == 2
    rows = list(csv.reader(result.paths[0].open(encoding='utf-8-sig', newline='')))
    assert rows[1] == ['张三', '00123', '包含,逗号\n和换行']


def test_csv_import_preserves_leading_zeros_and_long_ids(tmp_path):
    from factory.tables import convert_table
    from factory.common import Options
    source = tmp_path / 'ids.csv'
    source.write_text('编号,身份证,公式样文本\n00123,123456789012345678,=1+1\n', encoding='utf-8-sig')
    result = convert_table(source, 'xlsx', tmp_path, Options())
    book = load_workbook(result.paths[0])
    assert [c.value for c in book.active[2]] == ['00123', '123456789012345678', '=1+1']
    assert book.active['C2'].data_type == 's'
    book.close()


def test_formula_without_cache_is_reported(tmp_path):
    from factory.tables import convert_table
    from factory.common import Options
    book = Workbook()
    book.active.append(['值', '计算'])
    book.active.append([3, '=A2*2'])
    source = tmp_path / 'formula.xlsx'
    book.save(source)
    result = convert_table(source, 'csv', tmp_path, Options())
    assert any('公式' in warning for warning in result.warnings)


def test_excel_date_and_time_values(tmp_path):
    from datetime import date, datetime, time
    from factory.tables import convert_table
    from factory.common import Options
    book = Workbook()
    book.active.append([date(2026, 10, 17), time(9, 30), datetime(2026, 10, 17, 9, 30)])
    source = tmp_path / 'dates.xlsx'
    book.save(source)
    result = convert_table(source, 'csv', tmp_path, Options())
    row = next(csv.reader(result.paths[0].open(encoding='utf-8-sig')))
    assert row[1] == '09:30:00'
    assert '2026-10-17' in row[0] and '09:30:00' in row[2]


def test_semicolon_gb18030_csv(tmp_path):
    from factory.tables import convert_table
    from factory.common import Options
    source = tmp_path / '表.csv'
    source.write_text('名称;代码\n中文;0008\n', encoding='gb18030')
    result = convert_table(source, 'xlsx', tmp_path, Options(encoding='gb18030', delimiter=';'))
    book = load_workbook(result.paths[0])
    assert book.active['A2'].value == '中文'
    assert book.active['B2'].value == '0008'
    book.close()


def test_alpha_to_jpeg_is_white_and_warns(tmp_path):
    from factory.images import convert_image
    from factory.common import Options
    source = tmp_path / 'alpha.png'
    Image.new('RGBA', (30, 20), (255, 0, 0, 0)).save(source)
    result = convert_image(source, 'jpg', tmp_path, Options())
    with Image.open(result.paths[0]) as image:
        assert image.mode == 'RGB'
        assert min(image.getpixel((1, 1))) > 240
    assert result.warnings


def test_resize_preserves_ratio_and_exif_orientation(tmp_path):
    from factory.images import convert_image
    from factory.common import Options
    source = tmp_path / 'portrait.jpg'
    exif = Image.Exif()
    exif[274] = 6
    Image.new('RGB', (100, 50), 'red').save(source, exif=exif)
    result = convert_image(source, 'png', tmp_path, Options(max_size=40))
    with Image.open(result.paths[0]) as image:
        assert image.size == (20, 40)


def test_heic_real_round_trip(tmp_path):
    from factory.images import convert_image
    from factory.common import Options
    source = tmp_path / 'phone.png'
    Image.new('RGB', (64, 48), (20, 80, 140)).save(source)
    heic = convert_image(source, 'heic', tmp_path, Options()).paths[0]
    png = convert_image(heic, 'png', tmp_path, Options()).paths[0]
    with Image.open(png) as image:
        assert image.size == (64, 48)
    assert png != source


def test_corrupt_input_does_not_publish_partial_files(tmp_path):
    from factory.images import convert_image
    from factory.common import Options
    source = tmp_path / 'broken.heic'
    source.write_bytes(b'broken')
    output = tmp_path / 'out'
    with pytest.raises(Exception):
        convert_image(source, 'jpg', output, Options())
    assert not list(output.glob('*.jpg'))


@pytest.mark.parametrize('target', ['heic', 'png', 'webp', 'tiff'])
def test_default_conversion_removes_private_exif(tmp_path, target):
    from factory.images import convert_image
    from factory.common import Options
    source = tmp_path / 'private.jpg'
    exif = Image.Exif()
    exif[271] = 'Private Camera'
    exif[306] = '2025:01:02 03:04:05'
    Image.new('RGB', (64, 48), 'red').save(source, exif=exif)
    result = convert_image(source, target, tmp_path, Options())
    with Image.open(result.paths[0]) as image:
        assert 271 not in image.getexif() and 306 not in image.getexif()


def test_explicit_keep_metadata_retains_camera_info(tmp_path):
    from factory.images import convert_image
    from factory.common import Options
    source = tmp_path / 'private.jpg'
    exif = Image.Exif()
    exif[271] = 'My Camera'
    Image.new('RGB', (64, 48), 'red').save(source, exif=exif)
    result = convert_image(source, 'heic', tmp_path, Options(keep_metadata=True))
    with Image.open(result.paths[0]) as image:
        assert image.getexif()[271] == 'My Camera'
