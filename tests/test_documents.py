import pytest


def test_markdown_docx_round_trip(tmp_path):
    from factory.documents import convert_document
    from factory.common import Options
    from docx import Document
    source = tmp_path / 'example.md'
    source.write_text('# 中文标题\n\n这是 **加粗** 文本。\n\n|姓名|编号|\n|---|---|\n|张三|001|\n', encoding='utf-8')
    docx = convert_document(source, 'docx', tmp_path, Options()).paths[0]
    book = Document(docx)
    assert any('中文标题' in p.text for p in book.paragraphs)
    assert book.tables[0].cell(1, 0).text == '张三'
    md = convert_document(docx, 'md', tmp_path, Options()).paths[0]
    assert '中文标题' in md.read_text(encoding='utf-8')


def test_txt_pdf_round_trip_including_xml_symbols(tmp_path):
    from factory.documents import convert_document
    from factory.common import Options
    source = tmp_path / '原文.txt'
    source.write_text('你好，世界！\n公式 a < b & c > d\n第二行 αβ', encoding='utf-8')
    pdf = convert_document(source, 'pdf', tmp_path, Options(pdf_engine='basic')).paths[0]
    txt = convert_document(pdf, 'txt', tmp_path, Options()).paths[0]
    output = txt.read_text(encoding='utf-8')
    assert '你好，世界' in output
    assert 'a < b & c > d' in output
    assert 'αβ' in output
    assert source.read_text(encoding='utf-8').startswith('你好')


def test_pdf_to_editable_word(tmp_path):
    from factory.documents import convert_document
    from factory.common import Options
    from docx import Document
    source = tmp_path / 'hello.txt'
    source.write_text('Conversion test 123\n你好世界', encoding='utf-8')
    pdf = convert_document(source, 'pdf', tmp_path, Options(pdf_engine='basic')).paths[0]
    docx = convert_document(pdf, 'docx', tmp_path, Options()).paths[0]
    book = Document(docx)
    assert 'Conversion test' in '\n'.join(p.text for p in book.paragraphs)


def test_scan_pdf_without_ocr_has_actionable_error(tmp_path):
    import pymupdf
    from factory.documents import extract_pdf
    doc = pymupdf.open()
    doc.new_page()
    source = tmp_path / 'empty.pdf'
    doc.save(source)
    doc.close()
    with pytest.raises(ValueError, match='识别|文字'):
        extract_pdf(source, ocr=False)


def test_txt_to_markdown_does_not_interpret_headings(tmp_path):
    from factory.documents import convert_document
    from factory.common import Options
    source = tmp_path / 'literal.txt'
    source.write_text('# literal\n**not bold**\n', encoding='utf-8')
    result = convert_document(source, 'md', tmp_path, Options())
    assert '\\#' in result.paths[0].read_text(encoding='utf-8')


def test_relative_document_paths(tmp_path, monkeypatch):
    from pathlib import Path
    from factory.documents import convert_document
    from factory.common import Options
    monkeypatch.chdir(tmp_path)
    Path('input').mkdir()
    Path('input/中文.md').write_text('# 你好\n\n正文', encoding='utf-8')
    result = convert_document(Path('input/中文.md'), 'docx', Path('output'), Options())
    assert result.paths[0].is_file()


def test_word_markdown_keeps_images_with_spaces_in_name(tmp_path):
    from PIL import Image
    from docx import Document
    from factory.documents import convert_document
    from factory.common import Options
    image = tmp_path / 'picture.png'
    Image.new('RGB', (40, 40), 'blue').save(image)
    book = Document()
    book.add_paragraph('中文图片')
    book.add_picture(str(image))
    source = tmp_path / '中文 文档.docx'
    book.save(source)
    md = convert_document(source, 'md', tmp_path, Options()).paths[0]
    markdown = md.read_text(encoding='utf-8')
    assert '.factory-' not in markdown
    assert ':/' not in markdown, 'Image links should be portable relative paths'
    result = convert_document(md, 'docx', tmp_path, Options()).paths[0]
    assert len(Document(result).inline_shapes) == 1
